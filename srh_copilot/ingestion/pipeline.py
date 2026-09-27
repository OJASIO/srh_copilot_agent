"""Ingestion: data/raw/<agent>/<collection>/* -> chunks -> embeddings -> vector store.

Supported inputs: .md .txt .pdf .docx .pptx
Collection name = "<agent>/<subfolder>" which matches manifest.knowledge_collections.
Re-running replaces the collection, so the store never holds stale duplicates.
"""

from __future__ import annotations

import logging
import re
from pathlib import Path

from config.settings import PROJECT_ROOT, Settings
from core.services import ServiceContainer
from core.vector_store import Chunk

log = logging.getLogger(__name__)
RAW_DIR = PROJECT_ROOT / "data" / "raw"


# loaders

def load_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)
    if suffix == ".docx":
        from docx import Document

        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for r in t.rows:
                parts.append(" | ".join(c.text for c in r.cells))
        return "\n".join(parts)
    if suffix == ".pptx":
        from pptx import Presentation

        out = []
        for slide in Presentation(str(path)).slides:
            out += [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
        return "\n".join(out)
    raise ValueError(f"unsupported file type: {path.name}")


# chunking

def chunk_text(text: str, size: int, overlap: int) -> list[str]:
    """Paragraph-aware sliding window measured in characters."""
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    paras = [p.strip() for p in text.split("\n\n") if p.strip()]
    chunks, buf = [], ""
    for p in paras:
        if len(buf) + len(p) + 2 <= size:
            buf = f"{buf}\n\n{p}" if buf else p
            continue
        if buf:
            chunks.append(buf)
        while len(p) > size:  # very long paragraph: hard split
            chunks.append(p[:size])
            p = p[size - overlap:]
        buf = p
    if buf:
        chunks.append(buf)
    # add overlap from the previous chunk so context is not lost at boundaries
    out = []
    for i, c in enumerate(chunks):
        if i and overlap:
            c = chunks[i - 1][-overlap:] + "\n" + c
        out.append(c)
    return out


# pipeline

async def ingest_agent(services: ServiceContainer, agent_id: str, settings: Settings | None = None,
                       collections: list[str] | None = None) -> dict[str, int]:
    settings = settings or services.settings
    base = RAW_DIR / agent_id
    if not base.exists():
        raise FileNotFoundError(f"no raw data folder {base}")
    stats: dict[str, int] = {}
    for coll_dir in sorted(p for p in base.iterdir() if p.is_dir()):
        collection = f"{agent_id}/{coll_dir.name}"
        if collections and collection not in collections:
            continue
        removed = await services.vector_store.delete_collection(collection)
        chunks: list[Chunk] = []
        for f in sorted(coll_dir.rglob("*")):
            if not f.is_file() or f.name.startswith("."):
                continue
            try:
                text = load_text(f)
            except Exception as exc:
                log.warning("skip %s: %s", f, exc)
                continue
            rel = f.relative_to(RAW_DIR).as_posix()
            for i, piece in enumerate(chunk_text(text, settings.chunk_size, settings.chunk_overlap)):
                chunks.append(Chunk(text=piece, source=rel, collection=collection,
                                    metadata={"chunk": i, "file": f.name}))
        if chunks:
            batch = 64
            for i in range(0, len(chunks), batch):
                part = chunks[i:i + batch]
                embs = await services.embedder.embed([c.text for c in part])
                await services.vector_store.add(part, embs)
        stats[collection] = len(chunks)
        log.info("collection %s: removed %d, added %d chunks", collection, removed, len(chunks))
    return stats
