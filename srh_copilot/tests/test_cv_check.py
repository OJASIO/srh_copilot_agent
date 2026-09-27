"""CV Check tests. Offline: PDFs and DOCX files are generated on the fly and the
LLM is a stub that returns a fixed tiered review."""

from __future__ import annotations

import asyncio
import io
import json
import logging

import pytest

from agents.student_service.cv_check.anonymiser import anonymise_text
from agents.student_service.cv_check.ats_checker import check_ats
from agents.student_service.cv_check.extractor import extract_name_from_pdf, extract_text, fix_encoding
from agents.student_service.cv_check.language import detect_language
from agents.student_service.cv_check.report import render_markdown
from config.settings import Settings
from core.orchestrator import Orchestrator
from core.providers import BaseLLM
from core.registry import AgentRegistry
from core.schemas import AgentRequest, Attachment
from core.services import ServiceContainer

EN_CV = """Max Mustermann
max.mustermann@example.com | +49 151 12345678 | Heidelberg
linkedin.com/in/max-mustermann | github.com/maxm

PROFESSIONAL SUMMARY
Data science master student with three years of industry experience.

EXPERIENCE
Sep 2021 - Apr 2024  Developer, Infosys
- Built pipelines for 1M+ product records and reduced manual work by 30%

EDUCATION
04/2025 - present  M.Sc. Applied Data Science, SRH University Heidelberg

SKILLS
Python, SQL, Docker
"""

DE_CV = """Erika Musterfrau
Hauptstraße 12, 69117 Heidelberg
erika@example.de, 0176 1234567
Geburtsdatum: 01.01.1999

BERUFSERFAHRUNG
Seit 2022  Werkstudentin Datenanalyse bei der Beispiel GmbH

AUSBILDUNG
2019 - 2022  Bachelor Informatik, Universität Heidelberg

KENNTNISSE
Python und SQL, Deutsch Muttersprache, Englisch fließend
"""

REVIEW = {
    "overall_score": 7,
    "tier_1": [{"title": "Date formats", "detail": "Mixed 'Sep 2021' and '04/2025'.", "fix": "Use MM/YYYY throughout."}],
    "tier_2": [{"title": "Quantification", "detail": "Add numbers to the Infosys bullet."}],
    "tier_3": ["Career positioning"],
    "summary": "Solid CV with one consistency issue.",
    "ready": True,  # deliberately wrong: the task must derive it from tier_1
}


def make_pdf(text: str, name_size: int = 24) -> bytes:
    pymupdf = pytest.importorskip("pymupdf")
    doc = pymupdf.open()
    page = doc.new_page()
    lines = text.split("\n")
    page.insert_text((50, 60), lines[0], fontsize=name_size)
    y = 90
    for line in lines[1:]:
        page.insert_text((50, y), line, fontsize=10)
        y += 14
    out = doc.tobytes()
    doc.close()
    return out


def make_docx(text: str) -> bytes:
    from docx import Document

    doc = Document()
    for line in text.split("\n"):
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# extraction

def test_fix_encoding_handles_ligatures_and_custom_map():
    assert fix_encoding("ﬁnance") == "finance"
    assert fix_encoding("moƟon") == "motion"


def test_extract_text_pdf_and_docx():
    assert "Infosys" in extract_text("cv.pdf", make_pdf(EN_CV))
    assert "Infosys" in extract_text("cv.docx", make_docx(EN_CV))


def test_name_detection_uses_largest_font():
    assert extract_name_from_pdf(make_pdf(EN_CV)) == "Max Mustermann"


# language

def test_language_detection():
    assert detect_language(EN_CV) == "en"
    assert detect_language(DE_CV) == "de"
    assert detect_language("") == "unknown"


# anonymisation

def test_anonymiser_masks_pii_but_keeps_placeholders():
    clean = anonymise_text(DE_CV, filename="cv.docx")
    assert "erika@example.de" not in clean and "[EMAIL]" in clean
    assert "0176 1234567" not in clean and "[PHONE]" in clean
    assert "69117" not in clean and "[PLZ]" in clean
    assert "01.01.1999" not in clean and "[DOB REMOVED]" in clean
    assert "Hauptstraße 12" not in clean and "[ADDRESS REMOVED]" in clean
    assert "Erika Musterfrau" not in clean and "[PER]" in clean
    assert "Beispiel GmbH" in clean  # body content untouched


def test_anonymiser_uses_pdf_font_name_and_strips_links():
    clean = anonymise_text(EN_CV, filename="cv.pdf", data=make_pdf(EN_CV))
    assert "Max Mustermann" not in clean
    assert "[LINKEDIN REMOVED]" in clean and "[GITHUB REMOVED]" in clean


# ATS

def test_ats_checker_scores_and_issues():
    good = check_ats(DE_CV)
    assert good["score"] == 100 and good["issues"] == []
    bad = check_ats("Just some text with no contact data and ★ stars")
    kinds = {i["issue"].split(":")[0] for i in bad["issues"]}
    assert "No email address found" in kinds and "No phone number found" in kinds
    assert bad["score"] == 100 - 25 - 25 - 15 * 3 - 5
    mixed = check_ats(EN_CV)
    assert any("Inconsistent date formats" in i["issue"] for i in mixed["issues"])


# report

def test_render_markdown_german_labels():
    md = render_markdown({"language": "de", "language_label": "German (DE)",
                          "review": {**REVIEW, "ready": False}, "ats": {"score": 90, "issues": []}})
    assert "Kritische Fehler" in md and "Date formats" in md and "7/10" in md


# end to end through the orchestrator

class StubLLM(BaseLLM):
    def __init__(self):
        self.prompts: list[str] = []
        self.calls: list[dict] = []

    async def chat(self, messages, *, temperature=None, max_tokens=None, json_mode=False, json_schema=None) -> str:
        self.prompts.append(messages[-1]["content"])
        self.calls.append({"max_tokens": max_tokens, "json_mode": json_mode, "json_schema": json_schema})
        return "```json\n" + json.dumps(REVIEW) + "\n```"


@pytest.fixture()
def stack(tmp_path):
    settings = Settings(llm_provider="mock", vector_backend="memory", app_env="test",
                        vector_index_path=tmp_path / "index.json", api_key="test-key", enabled_agents=[])
    services = ServiceContainer.build(settings)
    services.llm = StubLLM()
    registry = AgentRegistry(services)
    registry.load()
    return services, registry


def _upload(stack, filename: str, data: bytes, message="Please review the attached document.", language="auto",
            inputs: dict | None = None):
    services, registry = stack
    orch = Orchestrator(registry, services)
    req = AgentRequest(session_id="s-cv", user_id="u1", agent_id="student_service", message=message,
                       language=language, inputs=inputs or {},
                       attachments=[Attachment(filename=filename, content_type="x", data=data)])
    return asyncio.run(orch.run(req))


def test_cv_check_end_to_end_pdf(stack):
    services, _ = stack
    resp = _upload(stack, "cv.pdf", make_pdf(EN_CV),
                   inputs={"job_description": "Data Analyst, requires SQL and Tableau. Contact jobs@example.com"})
    assert resp.task_id == "cv_check"
    findings = resp.structured["findings"]
    assert findings["language"] == "en"
    assert findings["review"]["overall_score"] == 7
    assert findings["review"]["ready"] is False  # derived from non-empty tier_1
    assert findings["ats"]["score"] < 100
    assert "Critical issues" in resp.content and "Date formats" in resp.content
    assert "not a legal, medical" not in resp.content  # no generic footer on CV reviews
    sent = services.llm.prompts[-1]
    assert "max.mustermann@example.com" not in sent and "Max Mustermann" not in sent
    assert "TARGET ROLE" in sent and "Tableau" in sent
    assert "jobs@example.com" not in sent  # recruiter contact data in the job description is masked
    assert "<cv_document>" in sent and "<job_description>" in sent
    assert findings["privacy"]["total"] >= 3 and findings["review_valid"] is True


def test_short_message_is_not_a_job_description(stack):
    services, _ = stack
    _upload(stack, "cv.pdf", make_pdf(EN_CV), message="please check my CV")
    assert "TARGET ROLE" not in services.llm.prompts[-1]


def test_long_pasted_message_is_the_job_description(stack):
    services, _ = stack
    advert = "We are hiring a Data Analyst (m/w/d). " + "You work with SQL, Tableau and Python every day. " * 5
    _upload(stack, "cv.pdf", make_pdf(EN_CV), message=advert)
    assert "TARGET ROLE" in services.llm.prompts[-1]


def test_cv_check_docx_german_prompt(stack):
    services, _ = stack
    resp = _upload(stack, "lebenslauf.docx", make_docx(DE_CV))
    assert resp.structured["findings"]["language"] == "de"
    assert "Kritische Fehler" in resp.content
    assert "PRÜFLISTE" in services.llm.prompts[-1]
    assert "ZIELSTELLE" not in services.llm.prompts[-1]  # default message is not a job description


def test_cv_check_language_override(stack):
    resp = _upload(stack, "lebenslauf.docx", make_docx(DE_CV), language="en")
    assert resp.structured["findings"]["language"] == "en"
    # the UI field wins over the API parameter
    resp = _upload(stack, "lebenslauf.docx", make_docx(DE_CV), language="en", inputs={"review_language": "de"})
    assert resp.structured["findings"]["language"] == "de"


def test_cv_check_rejects_unreadable_and_empty_files(stack):
    resp = _upload(stack, "cv.txt", b"hello")
    assert "PDF or DOCX" in resp.content
    resp = _upload(stack, "Lebenslauf_Erika_Musterfrau.pdf", b"not a pdf")
    assert "could not read" in resp.content
    resp = _upload(stack, "cv.docx", make_docx("too short"))
    assert "No usable text" in resp.content


def test_cv_check_survives_unparsable_llm_output(stack):
    services, _ = stack

    class Junk(BaseLLM):
        async def chat(self, messages, **kw):
            return "Sorry, I cannot help with that."

    services.llm = Junk()
    resp = _upload(stack, "cv.docx", make_docx(EN_CV))
    review = resp.structured["findings"]["review"]
    assert review["overall_score"] == 0 and review["ready"] is False
    assert resp.structured["findings"]["ats"]["score"] > 0  # ATS still runs
    assert resp.structured["findings"]["review_attempts"] == 2  # one retry
    assert resp.structured["findings"]["review_valid"] is False
    assert "could not be completed" in resp.content


def test_file_name_is_never_logged(stack, caplog):
    # file names usually contain the student's name
    with caplog.at_level(logging.DEBUG):
        _upload(stack, "Lebenslauf_Erika_Musterfrau.pdf", b"not a pdf")
    assert "Musterfrau" not in caplog.text


def test_follow_up_question_uses_the_earlier_review(stack):
    services, registry = stack
    orch = Orchestrator(registry, services)
    first = AgentRequest(session_id="s-follow", user_id="u1", agent_id="student_service", task_id="cv_check",
                         message="Please review the attached document.",
                         attachments=[Attachment(filename="cv.docx", content_type="x", data=make_docx(EN_CV))])
    asyncio.run(orch.run(first))
    follow = AgentRequest(session_id="s-follow", user_id="u1", agent_id="student_service", task_id="cv_check",
                          message="How do I fix the date format issue?")
    resp = asyncio.run(orch.run(follow))
    assert resp.structured.get("follow_up") is True
    assert "Please upload your CV" not in resp.content
    system = services.llm.calls  # StubLLM records every call
    assert len(system) == 2


def test_cv_without_file_and_without_review_asks_for_upload(stack):
    services, registry = stack
    resp = asyncio.run(Orchestrator(registry, services).run(AgentRequest(
        session_id="s-new", user_id="u1", agent_id="student_service", task_id="cv_check", message="check my CV")))
    assert "upload" in resp.content.lower()


def test_ats_english_month_names_are_one_format():
    cv = "a@b.de\n+49 151 1234567\nEDUCATION\nApril 2025 - Present\nEXPERIENCE\nSeptember 2021 - November 2024\nSKILLS"
    assert not any(i["code"] == "date_formats" for i in check_ats(cv)["issues"])


def test_ats_year_range_is_not_a_phone_number():
    issues = check_ats("a@b.de\nEDUCATION\n2017 - 2021 B.E.\nEXPERIENCE\nSKILLS")["issues"]
    assert any(i["code"] == "no_phone" for i in issues)


def test_ats_marketing_is_not_a_month_and_birth_date_is_ignored():
    cv = "a@b.de\n0176 1234567\nGeburtsdatum: 01.01.1999\nEDUCATION\n04/2025 - heute\nEXPERIENCE\nMarketing 2024\nSKILLS"
    assert not any(i["code"] == "date_formats" for i in check_ats(cv)["issues"])


def test_german_report_translates_ats_findings():
    md = render_markdown({"language": "de", "review": {**REVIEW, "ready": False},
                          "ats": check_ats("Nur Text ohne Kontakt")})
    assert "Keine E-Mail-Adresse gefunden" in md and "Standardabschnitt fehlt: Ausbildung" in md


def test_cut_off_review_is_retried_with_a_larger_budget(stack):
    from core.providers import LLMText

    services, _ = stack

    class CutOff(BaseLLM):
        def __init__(self):
            self.budgets = []

        async def chat(self, messages, *, max_tokens=None, **kw):
            self.budgets.append(max_tokens)
            text = LLMText(json.dumps(REVIEW)[:40] if len(self.budgets) == 1 else json.dumps(REVIEW))
            text.finish_reason = "length" if len(self.budgets) == 1 else "stop"
            return text

    services.llm = CutOff()
    resp = _upload(stack, "cv.docx", make_docx(EN_CV))
    assert services.llm.budgets == [2500, 4000]
    assert resp.structured["findings"]["review_valid"] and resp.structured["findings"]["review_attempts"] == 2
